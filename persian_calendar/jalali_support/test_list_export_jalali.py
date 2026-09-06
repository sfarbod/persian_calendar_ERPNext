# Copyright (c) 2026, Persian Calendar contributors
# License: MIT

"""v2.0.2 — List View Export Data (download_template / data_import.Exporter) Jalali tests."""

from __future__ import annotations

import json
from datetime import date, datetime
from io import BytesIO

import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils import cstr
from openpyxl import load_workbook

from persian_calendar.jalali_support.data_import_export import apply_data_import_export_patches
from persian_calendar.utils.data_io import convert_export_value


class TestListViewExportJalali(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		apply_data_import_export_patches()

	def setUp(self):
		super().setUp()
		frappe.local.response = frappe._dict()
		if getattr(frappe.local, "form_dict", None) is not None:
			frappe.local.form_dict.pop("export_dates_as_jalali", None)

	def _ensure_user(self) -> str:
		email = "jalali.list.export@test.local"
		if not frappe.db.exists("User", email):
			frappe.get_doc(
				{
					"doctype": "User",
					"email": email,
					"first_name": "Jalali List Export",
					"enabled": 1,
					"send_welcome_email": 0,
					"birth_date": "2026-09-06",
				}
			).insert(ignore_permissions=True)
		else:
			frappe.db.set_value("User", email, "birth_date", "2026-09-06", update_modified=False)
		return email

	@staticmethod
	def _xlsx_text(content: bytes) -> str:
		wb = load_workbook(BytesIO(content), data_only=True)
		ws = wb.active
		parts = []
		for row in ws.iter_rows(values_only=True):
			for cell in row:
				if cell is not None and str(cell).strip() != "":
					parts.append(str(cell))
		return "\n".join(parts)

	def _download(
		self,
		*,
		jalali: int,
		file_type: str = "CSV",
		export_records: str = "by_filter",
		via_form_dict: bool = False,
	) -> str:
		email = self._ensure_user()
		export_fields = {"User": ["name", "birth_date", "first_name"]}
		export_filters = [["User", "name", "=", email]]
		kwargs = {
			"doctype": "User",
			"export_fields": json.dumps(export_fields),
			"export_records": export_records,
			"export_filters": json.dumps(export_filters),
			"file_type": file_type,
		}
		from frappe.core.doctype.data_import.data_import import download_template

		if via_form_dict:
			frappe.local.form_dict = frappe._dict({**kwargs, "export_dates_as_jalali": str(jalali)})
			download_template(**kwargs)
		else:
			download_template(**kwargs, export_dates_as_jalali=jalali)

		if file_type == "CSV":
			return cstr(frappe.local.response.get("result") or "")
		content = frappe.local.response.get("filecontent")
		self.assertTrue(content)
		return self._xlsx_text(content)

	def test_date_off_stays_gregorian(self):
		result = self._download(jalali=0)
		self.assertIn("2026-09-06", result)
		self.assertNotIn("1405-06-15", result)

	def test_date_on_becomes_jalali(self):
		result = self._download(jalali=1)
		expected = convert_export_value("2026-09-06", "Date", True)
		self.assertEqual(expected, "1405-06-15")
		self.assertIn(expected, result)
		# Gregorian ISO must not appear as the exported birth_date cell
		self.assertNotIn('"2026-09-06"', result)

	def test_datetime_helpers_preserve_time(self):
		self.assertEqual(
			convert_export_value("2026-09-06 14:35:12", "Datetime", True),
			"1405-06-15 14:35:12",
		)
		self.assertEqual(
			convert_export_value("2026-09-06 14:35:12", "Datetime", False),
			"2026-09-06 14:35:12",
		)

	def test_empty_date_datetime_unchanged(self):
		self.assertEqual(convert_export_value(None, "Date", True), None)
		self.assertEqual(convert_export_value("", "Date", True), "")
		self.assertEqual(convert_export_value(None, "Datetime", True), None)
		self.assertEqual(convert_export_value("", "Datetime", True), "")

	def test_data_field_date_looking_string_not_converted(self):
		"""Fieldtype Data must not convert even if value looks like a date."""
		from frappe.core.doctype.data_import.exporter import Exporter

		exp = Exporter.__new__(Exporter)
		exp.export_dates_as_jalali = 1
		exp.doctype = "User"
		exp.fields = [
			frappe._dict(
				parent="User",
				fieldname="first_name",
				fieldtype="Data",
				is_child_table_field=False,
				child_table_df=None,
				hide_days=0,
			)
		]
		rows = Exporter.add_data_row(exp, "User", None, {"first_name": "2026-09-06"}, [], 0)
		self.assertEqual(rows[0][0], "2026-09-06")

	def test_exporter_converts_date_field_only(self):
		from frappe.core.doctype.data_import.exporter import Exporter

		exp = Exporter.__new__(Exporter)
		exp.export_dates_as_jalali = 1
		exp.doctype = "User"
		exp.fields = [
			frappe._dict(
				parent="User",
				fieldname="birth_date",
				fieldtype="Date",
				is_child_table_field=False,
				child_table_df=None,
				hide_days=0,
			),
			frappe._dict(
				parent="User",
				fieldname="first_name",
				fieldtype="Data",
				is_child_table_field=False,
				child_table_df=None,
				hide_days=0,
			),
		]
		rows = Exporter.add_data_row(
			exp,
			"User",
			None,
			{"birth_date": date(2026, 9, 6), "first_name": "2026-09-06"},
			[],
			0,
		)
		self.assertEqual(rows[0][0], "1405-06-15")
		self.assertEqual(rows[0][1], "2026-09-06")

	def test_child_table_date_and_datetime(self):
		from frappe.core.doctype.data_import.exporter import Exporter

		exp = Exporter.__new__(Exporter)
		exp.export_dates_as_jalali = 1
		exp.doctype = "Parent"
		child_df = frappe._dict(fieldname="items", label="Items")
		exp.fields = [
			frappe._dict(
				parent="Child",
				fieldname="schedule_date",
				fieldtype="Date",
				is_child_table_field=True,
				child_table_df=child_df,
				hide_days=0,
			),
			frappe._dict(
				parent="Child",
				fieldname="modified",
				fieldtype="Datetime",
				is_child_table_field=True,
				child_table_df=child_df,
				hide_days=0,
			),
		]
		rows = Exporter.add_data_row(
			exp,
			"Child",
			"items",
			{
				"schedule_date": "2026-09-06",
				"modified": datetime(2026, 9, 6, 14, 35, 12),
			},
			[],
			0,
		)
		self.assertEqual(rows[0][0], "1405-06-15")
		self.assertEqual(rows[0][1], "1405-06-15 14:35:12")

	def test_selected_and_filtered_records(self):
		# by_filter with name= is the List View "selected / filtered" path
		result = self._download(jalali=1, export_records="by_filter")
		self.assertIn("1405-06-15", result)
		self.assertIn(self._ensure_user(), result)

	def test_csv_and_excel(self):
		csv_out = self._download(jalali=1, file_type="CSV")
		xlsx_out = self._download(jalali=1, file_type="Excel")
		self.assertIn("1405-06-15", csv_out)
		self.assertIn("1405-06-15", xlsx_out)

	def test_form_dict_flag(self):
		result = self._download(jalali=1, via_form_dict=True)
		self.assertIn("1405-06-15", result)

	def test_unchecked_no_conversion_when_app_installed(self):
		result = self._download(jalali=0)
		self.assertIn("2026-09-06", result)
		self.assertNotIn("1405-06-15", result)

	def test_download_template_is_patched(self):
		from frappe.core.doctype.data_import import data_import as mod

		self.assertTrue(getattr(mod.download_template, "_jalali_patched", False))
